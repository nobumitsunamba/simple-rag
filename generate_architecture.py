"""Generate IT Architecture diagram as PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def add_box(slide, left, top, width, height, text, fill_color, font_size=9, bold=False, font_color=RGBColor(0xFF, 0xFF, 0xFF)):
    """Add a rounded rectangle box with text."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    shape.shadow.inherit = False

    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = font_color

    return shape


def add_arrow(slide, start_left, start_top, end_left, end_top):
    """Add a connector arrow."""
    connector = slide.shapes.add_connector(
        1,  # straight connector
        start_left, start_top, end_left, end_top
    )
    connector.line.color.rgb = RGBColor(0x71, 0x80, 0x96)
    connector.line.width = Pt(1.5)
    return connector


def add_label(slide, left, top, width, height, text, font_size=7):
    """Add a text label."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    return txBox


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Simple RAG - IT Architecture"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(12), Inches(0.4))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "AWS ECS Express Mode + Bedrock + Cognito + S3"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # === User Layer ===
    add_box(slide, Inches(0.5), Inches(1.8), Inches(1.8), Inches(1.2),
            "👤 ユーザー\n(ブラウザ)", RGBColor(0x4A, 0x55, 0x68), font_size=10, bold=True)

    # === ALB ===
    add_box(slide, Inches(3.0), Inches(1.8), Inches(1.8), Inches(1.2),
            "ALB\n(HTTPS/SSL)", RGBColor(0x8B, 0x5C, 0xF6), font_size=10, bold=True)

    # === ECS Container ===
    # Container background
    container_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                          Inches(5.5), Inches(1.2), Inches(4.5), Inches(5.5))
    container_bg.fill.solid()
    container_bg.fill.fore_color.rgb = RGBColor(0xEB, 0xF4, 0xFF)
    container_bg.line.color.rgb = RGBColor(0x66, 0x7E, 0xEA)
    container_bg.line.width = Pt(2)

    # Container label
    add_label(slide, Inches(5.7), Inches(1.3), Inches(4), Inches(0.4),
              "ECS Express Mode (Fargate) - 1 vCPU / 2GB", font_size=9)

    # FastAPI
    add_box(slide, Inches(5.8), Inches(1.8), Inches(3.9), Inches(0.7),
            "FastAPI (Python 3.11)", RGBColor(0x00, 0x96, 0x88), font_size=9, bold=True)

    # Internal modules
    add_box(slide, Inches(5.8), Inches(2.7), Inches(1.8), Inches(0.8),
            "📄 Document\nProcessor", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    add_box(slide, Inches(7.9), Inches(2.7), Inches(1.8), Inches(0.8),
            "🔍 Vector Store\n(Hybrid Search)", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    add_box(slide, Inches(5.8), Inches(3.7), Inches(1.8), Inches(0.8),
            "🤖 RAG Engine\n(Rerank + Generate)", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    add_box(slide, Inches(7.9), Inches(3.7), Inches(1.8), Inches(0.8),
            "🔐 Auth Module\n(Cognito)", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    add_box(slide, Inches(5.8), Inches(4.7), Inches(1.8), Inches(0.8),
            "📊 Audit Log", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    add_box(slide, Inches(7.9), Inches(4.7), Inches(1.8), Inches(0.8),
            "💾 Knowledge\nStore (S3)", RGBColor(0x66, 0x7E, 0xEA), font_size=8)

    # === External AWS Services ===
    # Bedrock
    add_box(slide, Inches(10.8), Inches(1.5), Inches(2.0), Inches(1.0),
            "Amazon Bedrock\nTitan Embeddings V2", RGBColor(0xE5, 0x3E, 0x3E), font_size=8, bold=True)

    # Anthropic
    add_box(slide, Inches(10.8), Inches(2.8), Inches(2.0), Inches(1.0),
            "Anthropic API\nClaude Sonnet 4.6", RGBColor(0xD6, 0x9E, 0x2E), font_size=8, bold=True)

    # S3
    add_box(slide, Inches(10.8), Inches(4.1), Inches(2.0), Inches(1.0),
            "Amazon S3\nナレッジ永続化\n+ ファイル保存", RGBColor(0x38, 0xA1, 0x69), font_size=8, bold=True)

    # Cognito
    add_box(slide, Inches(10.8), Inches(5.4), Inches(2.0), Inches(1.0),
            "Amazon Cognito\nUser Pool", RGBColor(0xDD, 0x6B, 0x20), font_size=8, bold=True)

    # CloudWatch
    add_box(slide, Inches(10.8), Inches(6.7), Inches(2.0), Inches(0.8),
            "CloudWatch Logs\n監査ログ", RGBColor(0x55, 0x55, 0x55), font_size=8, bold=True)

    # === CI/CD ===
    add_box(slide, Inches(0.5), Inches(5.5), Inches(1.8), Inches(1.0),
            "GitHub\nActions", RGBColor(0x24, 0x29, 0x2E), font_size=9, bold=True)

    add_box(slide, Inches(3.0), Inches(5.5), Inches(1.8), Inches(1.0),
            "Amazon ECR\n(Container Registry)", RGBColor(0xE5, 0x3E, 0x3E), font_size=8, bold=True)

    # === Arrows (simplified as labels) ===
    add_label(slide, Inches(2.3), Inches(2.1), Inches(0.7), Inches(0.4), "→", font_size=16)
    add_label(slide, Inches(4.8), Inches(2.1), Inches(0.7), Inches(0.4), "→", font_size=16)
    add_label(slide, Inches(9.8), Inches(1.7), Inches(1.0), Inches(0.4), "→", font_size=14)
    add_label(slide, Inches(9.8), Inches(3.0), Inches(1.0), Inches(0.4), "→", font_size=14)
    add_label(slide, Inches(9.8), Inches(4.3), Inches(1.0), Inches(0.4), "→", font_size=14)
    add_label(slide, Inches(9.8), Inches(5.6), Inches(1.0), Inches(0.4), "→", font_size=14)
    add_label(slide, Inches(9.8), Inches(6.8), Inches(1.0), Inches(0.4), "→", font_size=14)

    # CI/CD arrows
    add_label(slide, Inches(2.3), Inches(5.7), Inches(0.7), Inches(0.4), "→", font_size=16)
    add_label(slide, Inches(4.8), Inches(5.7), Inches(0.7), Inches(0.4), "→", font_size=14)

    # CI/CD label
    add_label(slide, Inches(0.5), Inches(5.0), Inches(4.5), Inches(0.4),
              "CI/CD Pipeline (git push → build → deploy)", font_size=8)

    # Legend
    add_label(slide, Inches(0.5), Inches(6.8), Inches(4.5), Inches(0.5),
              "Auto Scaling: 0〜20 tasks | Scale-to-Zero 対応 | OIDC認証 (キーレス)", font_size=8)

    # Save
    output_path = "docs/architecture.pptx"
    prs.save(output_path)
    print(f"アーキテクチャ図を保存しました: {output_path}")


if __name__ == "__main__":
    main()
